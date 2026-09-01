"""Build results/summary_grid.xlsx from results/grid_summary.csv (+ per-variant descriptions).
usage: /usr/bin/python3 make_excel.py   (system python has openpyxl)"""
import csv, os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "results", "summary_grid.xlsx")

DESC = {
    ("BGN", "baseline"): "BGN 1999 exactly as published (paper Table I calibration), rank chars, unpenalized market.",
    ("BGN", "crash"): "Rare-disaster kernel (lognormal x kappa^D), disaster-exposure project types, priced survival, compensation omega=1.3, level features.",
    ("BGN", "regime-g"): "2-state observable regime multiplying the price of risk gamma by [0.5, 2.0]; closed-form two-basis valuation; regime exported as conditioning feature.",
    ("BGN", "regime-g wide"): "Same regime-gamma design with wide multipliers [0.3, 3.0]; regime-shifting char levels break rolling FMR/FF.",
    ("BGN", "gamma(r) nonlin"): "gamma = smooth logistic function of the interest rate r (continuous observable state, 0.5-3.0); operator-chain pricing on an r-grid.",
    ("KP", "baseline"): "KP14 as published (paper Table II calibration; r=5% as in Code/), rank chars, unpenalized market.",
    ("KP", "crash"): "Disaster recipe ported to KP14: kernel x kappa^D, exposure-typed project kill fractions, priced survival, compensation.",
    ("KP", "regime-g"): "2-state regime multiplying both prices of risk (gamma_x, gamma_z) by [0.5, 2.0]; coupled A-coefficients and 4-state G tables.",
    ("KP", "gamma(y) common"): "gamma multiplier = smooth function of a continuous OU state y, common to both shocks (21-node y-grid, GH quadrature mixing).",
    ("KP", "gamma(y) rotation"): "gamma_x varies with y (0.85-3.0) while gamma_z stays fixed - a rotation of the price-of-risk vector, not a rescaling.",
    ("KP", "regime uneven"): "2-state regime scaling gamma_x only (x-channel swing), leaving gamma_z fixed - uneven version of regime-g.",
    ("KP", "exposure-types x regime"): "Firm types load cash flows as x^beta (beta 1.0/1.8/3.0, GBM powers closed-form) crossed with regime-gamma [0.5,2] + calm-value compensation 1.2. First KP room; not harvested (exposure leaks into raw char levels, rolling FMR conditions on them).",
    ("KP", "priced-vol exposures (vy)"): "Priced stationary OU factor y (kappa=0.35, gamma_v=1.2); types load e^{beta y} (beta 0.02/0.06/0.12, premia 2-12%/yr) + y=0 compensation 1.2. Levels stationary by design. Largest room of the project (+0.28); harvested: rff_ens beats FMR by +0.026 (t=5.2).",
    ("GS", "baseline"): "GS21 with the exact-kernel Python re-solve (row-renormalized SDF, converged fixed point), constant gamma_x=0.5.",
    ("GS", "crash"): "Disaster recipe in GS (re-based on exact kernel): kernel x kappa^D, capital-destruction types, endogenous deleveraging and real defaults.",
    ("GS", "gamma(x)"): "State-dependent price of risk gamma(x) = clip(0.5 - 0.28 x/sd, 0.05, 1), re-solved value functions; conditioning/timing channel only.",
    ("GS", "gamma(x) nonlin"): "Logistic (strongly nonlinear) gamma(x); still zero cross-sectional room - direction-invariance theorem.",
    ("GS", "regime-g"): "2-regime gamma multiplier [0.6, 3.0] with coupled 2-regime solver; timing gap without cross-sectional room.",
    ("GS", "exposure-types x regime"): "Firm types with heterogeneous production exposure e^{beta x} (3 types) crossed with regime-gamma [0.6,3.0].",
    ("GS", "exposure-types comp."): "Same exposure types + calm-value compensation (a-shifts dosed from measured calm log-value gaps) so values do not reveal types monotonically.",
    ("GS", "exposure-types wide (bx7)"): "5 types, beta 1-7, compensation dosed by value-gap slope, leverage characteristic added. Room +0.018, harvested vs classical methods: rff beats FMR by +0.016 (t=10.4); linear ranks+levels ridge comes within 0.001 of DKKM. FMR overflows in 0.8% of months (extreme raw levels).",
}

COLS = ["model", "variant", "SR_max", "lin_ceil", "nl_ceil", "room", "FMR", "FF", "lin", "RFF", "gap_RFF-lin", "t_vs_FMR"]
COLDEF = [
    ("model", "Structural model: BGN (Berk-Green-Naik 1999), KP (Kogan-Papanikolaou 2014), GS (Gomes-Schmid 2021)."),
    ("variant", "Economy variant; see the description column and REPORT.md."),
    ("SR_max", "Mean conditional maximum Sharpe ratio of the economy (population, monthly)."),
    ("lin_ceil", "Best population constant-theta ceiling over linear bases (rank chars, +regime feature, +levels)."),
    ("nl_ceil", "Best population constant-theta ceiling over nonlinear bases (RFF up to P=3600, poly2, bins)."),
    ("room", "nl_ceil - lin_ceil: population Sharpe available to nonlinear methods beyond any linear method."),
    ("FMR", "Best realized mean conditional SR, Fama-MacBeth regression portfolios (raw chars, rolling 360m MVE)."),
    ("FF", "Best realized SR, Fama-French sorted factor portfolios (rolling 360m MVE)."),
    ("lin", "Best realized SR over linear methods: ridge on rank chars (linrank) or ranks+levels (linlev)."),
    ("RFF", "Best realized SR over DKKM variants: random-Fourier-feature ridge (rff, rff_ens, rff_lev, rff_lev_ens), kappa grid to 10."),
    ("gap_RFF-lin", "RFF - lin: realized nonlinear-over-linear gap."),
    ("t_vs_FMR", "Paired monthly t-statistic of the best RFF variant's conditional SR against FMR."),
]

with open(os.path.join(HERE, "results", "grid_summary.csv")) as f:
    rows = list(csv.DictReader(f))

wb = Workbook()
ws = wb.active
ws.title = "Summary grid"
hdr_fill = PatternFill("solid", fgColor="1F4E79")
hdr_font = Font(bold=True, color="FFFFFF")
thin = Border(bottom=Side(style="thin", color="D0D0D0"))
model_fill = {"BGN": PatternFill("solid", fgColor="EAF1FB"),
              "KP": PatternFill("solid", fgColor="FDF2E3"),
              "GS": PatternFill("solid", fgColor="EAF7EA")}

headers = COLS + ["description"]
ws.append(headers)
for c, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=c)
    cell.fill, cell.font = hdr_fill, hdr_font
    cell.alignment = Alignment(horizontal="center")

for r in rows:
    vals = []
    for c in COLS:
        v = r.get(c, "")
        try:
            v = float(v)
        except (TypeError, ValueError):
            pass
        vals.append(v)
    vals.append(DESC.get((r["model"], r["variant"]), ""))
    ws.append(vals)

for i in range(2, ws.max_row + 1):
    m = ws.cell(row=i, column=1).value
    for j in range(1, len(headers) + 1):
        cell = ws.cell(row=i, column=j)
        cell.border = thin
        if m in model_fill:
            cell.fill = model_fill[m]
        if j >= 3 and j <= len(COLS) and isinstance(cell.value, float):
            cell.number_format = "0.000"
    ws.cell(row=i, column=len(headers)).alignment = Alignment(wrap_text=True, vertical="top")

widths = [7, 26, 8, 8, 8, 8, 7, 7, 7, 7, 11, 9, 95]
for j, w in enumerate(widths, 1):
    ws.column_dimensions[ws.cell(row=1, column=j).column_letter].width = w
ws.freeze_panes = "C2"

ws2 = wb.create_sheet("Definitions")
ws2.append(["column", "definition"])
for c in (1, 2):
    cell = ws2.cell(row=1, column=c)
    cell.fill, cell.font = hdr_fill, hdr_font
for name, d in COLDEF:
    ws2.append([name, d])
ws2.column_dimensions["A"].width = 14
ws2.column_dimensions["B"].width = 110
for i in range(2, ws2.max_row + 1):
    ws2.cell(row=i, column=2).alignment = Alignment(wrap_text=True)

wb.save(OUT)
print(f"wrote {OUT}: {ws.max_row - 1} variants")
